"""Control crosswalks — DISA CCI to NIST SP 800-53, and NIST 800-53 to ISO 27001.

Two authoritative crosswalks ship with the project artefacts:

* ``U_CCI_List.xml`` (DISA) maps each CCI to the NIST SP 800-53 elements it
  implements. Every DISA STIG rule cites CCIs, so this is what lets a STIG
  finding carry real 800-53 references instead of a guessed mapping.
* The NIST OLIR workbook maps SP 800-53 Rev 5 controls to ISO/IEC 27001:2022
  clauses and Annex A controls.

Control identifiers are normalised to a single canonical form so the two
vocabularies can be joined: uppercase family, zero-padded two-digit number, and
zero-padded parenthesised enhancement — ``AC-17(02)``. Statement letters
(``AC-1 a 1``) are dropped because they address sub-paragraphs of the same
control.
"""

from __future__ import annotations

import json
import re
import xml.etree.ElementTree as ElementTree
from pathlib import Path

from backend.app.config import FILE_ENCODING, FRAMEWORKS_DIR

CCI_NS = {"c": "http://iase.disa.mil/cci"}
CROSSWALK_DIR = FRAMEWORKS_DIR / "crosswalk"
_NIST_REV5_TITLE = "NIST SP 800-53 Revision 5"

# "AC-17 (2) (a)" / "ac-17(2)" / "AC-1 a 1" -> family, number, optional enhancement
_CONTROL_RE = re.compile(
    r"^\s*(?P<family>[A-Za-z]{2,3})\s*-\s*(?P<number>\d{1,3})"
    r"(?:\s*\(\s*(?P<enhancement>\d{1,3})\s*\))?"
)


def normalise_control_id(raw: str) -> str | None:
    """Normalise a NIST SP 800-53 control identifier to ``AC-17(02)`` form.

    Returns None when the string does not name a control (OLIR rows also carry
    bare ISO clause numbers such as ``5.2``).
    """
    match = _CONTROL_RE.match(raw or "")
    if not match:
        return None
    family = match.group("family").upper()
    number = int(match.group("number"))
    enhancement = match.group("enhancement")
    base = f"{family}-{number:02d}"
    return f"{base}({int(enhancement):02d})" if enhancement else base


def build_cci_to_nist(cci_xml: Path) -> dict[str, list[str]]:
    """Build the CCI to NIST SP 800-53 Rev 5 crosswalk from the DISA CCI list.

    Falls back to the most recent revision present for a CCI when Rev 5 has no
    reference, so older CCIs still resolve instead of silently mapping to
    nothing.
    """
    root = ElementTree.parse(str(cci_xml)).getroot()
    mapping: dict[str, list[str]] = {}

    for item in root.findall(".//c:cci_item", CCI_NS):
        cci_id = item.get("id")
        if not cci_id:
            continue
        by_revision: dict[str, set[str]] = {}
        for reference in item.findall("./c:references/c:reference", CCI_NS):
            title = reference.get("title") or ""
            if not title.startswith("NIST SP 800-53"):
                continue
            if title.startswith("NIST SP 800-53A"):
                continue
            control = normalise_control_id(reference.get("index") or "")
            if control:
                by_revision.setdefault(title, set()).add(control)

        controls = by_revision.get(_NIST_REV5_TITLE)
        if not controls:
            for title in sorted(by_revision, reverse=True):
                controls = by_revision[title]
                break
        if controls:
            mapping[cci_id] = sorted(controls)

    return mapping


def build_nist_to_iso(olir_xlsx: Path) -> dict[str, list[str]]:
    """Build the NIST SP 800-53 Rev 5 to ISO/IEC 27001:2022 crosswalk.

    Reads every relationship worksheet in the OLIR workbook. Rows whose focal
    element is not a control identifier are skipped.

    Raises:
        ArtifactMissingError: When openpyxl is not installed.
    """
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - dependency is installed
        from backend.core_errors import ArtifactMissingError

        raise ArtifactMissingError(
            "openpyxl is required to read the NIST/ISO OLIR workbook. "
            "Run Step 1 in MANUAL_COMMANDS.md, then re-run this command."
        ) from exc

    workbook = openpyxl.load_workbook(str(olir_xlsx), read_only=True, data_only=True)
    mapping: dict[str, set[str]] = {}

    for sheet_name in workbook.sheetnames:
        if sheet_name.lower() == "definitions":
            continue
        sheet = workbook[sheet_name]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row or len(row) < 4:
                continue
            control = normalise_control_id(str(row[0] or ""))
            iso_element = str(row[3] or "").strip()
            if control and iso_element:
                mapping.setdefault(control, set()).add(iso_element)

    workbook.close()
    return {control: sorted(elements) for control, elements in sorted(mapping.items())}


def write_crosswalk(name: str, mapping: dict[str, list[str]]) -> Path:
    """Persist a crosswalk as deterministic sorted JSON."""
    CROSSWALK_DIR.mkdir(parents=True, exist_ok=True)
    out = CROSSWALK_DIR / f"{name}.json"
    out.write_text(
        json.dumps(mapping, indent=2, sort_keys=True) + "\n", encoding=FILE_ENCODING
    )
    return out


def load_crosswalk(name: str) -> dict[str, list[str]]:
    """Load a built crosswalk, returning {} when it has not been built yet."""
    path = CROSSWALK_DIR / f"{name}.json"
    if not path.exists():
        return {}
    loaded: dict[str, list[str]] = json.loads(path.read_text(encoding=FILE_ENCODING))
    return loaded


def resolve_nist_from_cci(cci_ids: list[str]) -> list[str]:
    """Map a rule's CCI identifiers to NIST SP 800-53 controls."""
    crosswalk = load_crosswalk("cci_to_nist_800_53")
    resolved: set[str] = set()
    for cci in cci_ids:
        resolved.update(crosswalk.get(cci, []))
    return sorted(resolved)


def resolve_iso_from_nist(nist_ids: list[str]) -> list[str]:
    """Map NIST SP 800-53 controls to ISO/IEC 27001:2022 elements."""
    crosswalk = load_crosswalk("nist_800_53_to_iso_27001")
    resolved: set[str] = set()
    for control in nist_ids:
        normalised = normalise_control_id(control) or control
        resolved.update(crosswalk.get(normalised, []))
    return sorted(resolved)
